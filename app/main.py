import json
import os
import re
import subprocess
import threading
import winreg
from contextlib import asynccontextmanager
from datetime import date, datetime
from pathlib import Path

import markdown as md_lib
from fastapi import FastAPI, Depends, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import text
from sqlalchemy.orm import Session

from .database import Base, engine, get_db, SessionLocal
from .excel_export import gerar_excel
from .models import Documento, Revisao
from .seed import popular_banco

TEMPLATES_DIR    = Path(__file__).resolve().parent.parent / "templates"
STATIC_DIR       = Path(__file__).resolve().parent.parent / "static"
FINDABC_DIR      = Path("D:/02_Finaud/Projetos/ativos/findabc")
WORD_TEMP_DIR    = Path(__file__).resolve().parent.parent / "word_temp"
AUDIT_REPORT_DIR = Path(r"D:\02_Finaud\Projetos\ativos\_auditoria_seguranca\relatorios")
AUDIT_SCRIPT     = Path(r"D:\02_Finaud\Projetos\ativos\_auditoria_seguranca\security_audit.py")
AUDIT_PYTHON     = Path(r"D:\02_Finaud\Projetos\ativos\_auditoria_seguranca\.venv\Scripts\python.exe")
LOGO_PATH        = STATIC_DIR / "img" / "logo_finaud.png"

# ── Estado da auditoria em background ────────────────────────────────────────
_audit_lock   = threading.Lock()
_audit_status: dict = {"running": False, "started_at": None, "finished_at": None, "error": None}

def _run_audit_bg() -> None:
    global _audit_status
    try:
        subprocess.run(
            [str(AUDIT_PYTHON), str(AUDIT_SCRIPT)],
            capture_output=True, timeout=900,
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
        with _audit_lock:
            _audit_status["error"] = None
    except Exception as e:
        with _audit_lock:
            _audit_status["error"] = str(e)
    finally:
        with _audit_lock:
            _audit_status["running"]     = False
            _audit_status["finished_at"] = datetime.now().strftime("%d/%m/%Y %H:%M")


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(bind=engine)
    try:
        with engine.connect() as conn:
            conn.execute(text("ALTER TABLE revisoes ADD COLUMN hora TEXT"))
            conn.commit()
    except Exception:  # nosec B110 — coluna já existe no banco (ALTER TABLE idempotente)
        pass  # coluna já existe
    with SessionLocal() as db:
        popular_banco(db)
    yield


app = FastAPI(title="Painel de Conformidade — Finaud", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
if AUDIT_REPORT_DIR.exists():
    app.mount("/relatorios", StaticFiles(directory=AUDIT_REPORT_DIR), name="relatorios")
templates = Jinja2Templates(directory=TEMPLATES_DIR)


# ── utilitários ──────────────────────────────────────────────────────────────

def _ler_arquivo(arquivo: str) -> str | None:
    """Lê .md ou .html relativo ao findabc. Retorna None se não encontrar."""
    path = FINDABC_DIR / arquivo
    if not path.exists():
        return None
    return path.read_text(encoding="utf-8")


def _renderizar(arquivo: str) -> str:
    """Converte .md para HTML. Para outros formatos retorna texto cru."""
    conteudo = _ler_arquivo(arquivo)
    if conteudo is None:
        return "<p><em>Arquivo não encontrado.</em></p>"
    if arquivo.endswith(".md"):
        return md_lib.markdown(conteudo, extensions=["tables", "fenced_code"])
    return f"<pre>{conteudo}</pre>"


def _caminho_word() -> str | None:
    """Encontra o executável do Word via registro do Windows."""
    try:
        with winreg.OpenKey(
            winreg.HKEY_LOCAL_MACHINE,
            r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\WINWORD.EXE",
        ) as key:
            path = winreg.QueryValue(key, None)
            if path and Path(path).exists():
                return path
    except Exception:
        pass
    return None


# ── Conversão Markdown ↔ Word (.docx) ────────────────────────────────────────

def _is_separador_tabela(linha: str) -> bool:
    return bool(re.match(r'^\|(\s*:?-+:?\s*\|)+$', linha.strip()))


def _add_md_runs(paragrafo, texto: str) -> None:
    """Adiciona texto formatado (bold/italic markdown) a um parágrafo Word."""
    partes = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', texto)
    for parte in partes:
        if parte.startswith('**') and parte.endswith('**'):
            paragrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith('*') and parte.endswith('*'):
            paragrafo.add_run(parte[1:-1]).italic = True
        elif parte.startswith('`') and parte.endswith('`'):
            run = paragrafo.add_run(parte[1:-1])
            run.font.name = 'Courier New'
        elif parte:
            paragrafo.add_run(parte)


def _strip_md_inline(texto: str) -> str:
    texto = re.sub(r'\*\*([^*]+)\*\*', r'\1', texto)
    texto = re.sub(r'\*([^*]+)\*', r'\1', texto)
    texto = re.sub(r'`([^`]+)`', r'\1', texto)
    return texto


def _md_para_docx(conteudo: str, saida: Path) -> None:
    """Converte markdown para .docx usando python-docx."""
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3)
        section.right_margin  = Cm(2.5)

    linhas = conteudo.split('\n')
    i = 0
    while i < len(linhas):
        linha = linhas[i]

        if linha.startswith('### '):
            doc.add_heading(linha[4:].strip(), level=3)
        elif linha.startswith('## '):
            doc.add_heading(linha[3:].strip(), level=2)
        elif linha.startswith('# '):
            doc.add_heading(linha[2:].strip(), level=1)
        elif linha.strip() in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(4)
            run = p.add_run('─' * 60)
            run.font.color.rgb = None
        elif linha.strip().startswith('|') and '|' in linha[1:]:
            # Coleta linhas da tabela
            linhas_tabela: list[list[str]] = []
            while i < len(linhas) and linhas[i].strip().startswith('|') and '|' in linhas[i][1:]:
                if not _is_separador_tabela(linhas[i]):
                    celulas = [c.strip() for c in linhas[i].strip().strip('|').split('|')]
                    linhas_tabela.append(celulas)
                i += 1

            if linhas_tabela:
                n_cols = max(len(r) for r in linhas_tabela)
                tabela = doc.add_table(rows=len(linhas_tabela), cols=n_cols)
                tabela.style = 'Table Grid'
                for r_idx, linha_t in enumerate(linhas_tabela):
                    for c_idx in range(n_cols):
                        cell_text = linha_t[c_idx] if c_idx < len(linha_t) else ''
                        cell = tabela.cell(r_idx, c_idx)
                        cell.text = _strip_md_inline(cell_text)
                        if r_idx == 0:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True
            continue  # i já avançado pelo while interno
        elif linha.startswith('- ') or linha.startswith('* '):
            doc.add_paragraph(_strip_md_inline(linha[2:].strip()), style='List Bullet')
        elif re.match(r'^\d+\. ', linha):
            doc.add_paragraph(_strip_md_inline(re.sub(r'^\d+\. ', '', linha)), style='List Number')
        elif linha.strip():
            p = doc.add_paragraph()
            _add_md_runs(p, linha)
        # else: linha vazia — pula

        i += 1

    doc.save(str(saida))


def _runs_para_md(paragrafo) -> str:
    md = ''
    for run in paragrafo.runs:
        t = run.text
        if not t:
            continue
        if run.bold and run.italic:
            md += f'***{t}***'
        elif run.bold:
            md += f'**{t}**'
        elif run.italic:
            md += f'*{t}*'
        else:
            md += t
    return md


def _docx_para_md(docx_path: Path) -> str:
    """Converte .docx editado no Word de volta para markdown."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(docx_path))
    linhas: list[str] = []

    for bloco in doc.element.body:
        tag = bloco.tag.split('}')[-1] if '}' in bloco.tag else bloco.tag

        if tag == 'tbl':
            tabela = Table(bloco, doc)
            for r_idx, linha_t in enumerate(tabela.rows):
                celulas = [c.text.replace('\n', ' ').strip() for c in linha_t.cells]
                linhas.append('| ' + ' | '.join(celulas) + ' |')
                if r_idx == 0:
                    linhas.append('| ' + ' | '.join(['---'] * len(celulas)) + ' |')
            linhas.append('')

        elif tag == 'p':
            para = Paragraph(bloco, doc)
            estilo = para.style.name if para.style else 'Normal'
            texto  = para.text.strip()

            if 'Heading 1' in estilo:
                if texto:
                    linhas.append(f'# {texto}')
            elif 'Heading 2' in estilo:
                if texto:
                    linhas.append(f'## {texto}')
            elif 'Heading 3' in estilo:
                if texto:
                    linhas.append(f'### {texto}')
            elif 'List Bullet' in estilo:
                md = _runs_para_md(para)
                if md.strip():
                    linhas.append(f'- {md.strip()}')
            elif 'List Number' in estilo:
                md = _runs_para_md(para)
                if md.strip():
                    linhas.append(f'1. {md.strip()}')
            elif texto and set(texto) <= {'─', '—', '-', ' '}:
                linhas.append('---')
            else:
                md = _runs_para_md(para)
                linhas.append(md if md else '')

    resultado = '\n'.join(linhas)
    resultado = re.sub(r'\n{3,}', '\n\n', resultado)
    return resultado.strip() + '\n'


# ── health ───────────────────────────────────────────────────────────────────

@app.get("/health")
def health():
    return {"status": "ok"}


# ── dashboard ────────────────────────────────────────────────────────────────

_CATEGORIAS_ORDEM = ["Política", "Procedimento", "Governança", "Técnico", "Certificado"]


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, db: Session = Depends(get_db)):
    docs = db.query(Documento).order_by(Documento.nome).all()

    por_categoria: dict[str, list] = {}
    for cat in _CATEGORIAS_ORDEM:
        cat_docs = [d for d in docs if d.categoria == cat]
        if cat_docs:
            por_categoria[cat] = cat_docs
    for d in docs:
        if d.categoria not in por_categoria:
            por_categoria.setdefault(d.categoria, []).append(d)

    docs_json = json.dumps([{
        "id":           d.id,
        "nome":         d.nome,
        "categoria":    d.categoria,
        "periodicidade":d.periodicidade,
        "versao":       d.versao_atual,
        "proxima_versao": d.proxima_versao_sugerida,
        "prazo_str":    d.proxima_revisao_str,
        "dias":         d.dias_para_revisao,
        "status":       d.status,
        "tem_arquivo":  bool(d.arquivo and d.categoria != "Certificado"),
    } for d in docs], ensure_ascii=False)

    return templates.TemplateResponse(request, "dashboard.html", {
        "request":       request,
        "por_categoria": por_categoria,
        "docs_json":     docs_json,
        "total":         len(docs),
    })


# ── ver / exportar documento ─────────────────────────────────────────────────

@app.get("/ver/{doc_id}", response_class=HTMLResponse)
def ver_documento(doc_id: int, request: Request, db: Session = Depends(get_db)):
    doc = db.query(Documento).filter(Documento.id == doc_id).first()
    if not doc:
        return HTMLResponse("Documento não encontrado.", status_code=404)

    eh_pdf = doc.arquivo and doc.arquivo.endswith(".pdf")
    conteudo_html = "" if eh_pdf else _renderizar(doc.arquivo or "")

    return templates.TemplateResponse(request, "ver.html", {
        "request":       request,
        "doc":           doc,
        "conteudo_html": conteudo_html,
        "eh_pdf":        eh_pdf,
    })


# ── abrir arquivo no editor ──────────────────────────────────────────────────

@app.get("/abrir/{doc_id}")
def abrir_arquivo(doc_id: int, ajax: bool = False, db: Session = Depends(get_db)):
    doc = db.query(Documento).filter(Documento.id == doc_id).first()
    if doc and doc.arquivo:
        md_path = FINDABC_DIR / doc.arquivo
        if md_path.exists() and doc.arquivo.endswith(".md") and doc.categoria != "Certificado":
            WORD_TEMP_DIR.mkdir(exist_ok=True)
            docx_path = WORD_TEMP_DIR / f"{doc_id}.docx"
            # Só regenera se ainda não existe (preserva edição pendente)
            if not docx_path.exists():
                conteudo = md_path.read_text(encoding="utf-8")
                _md_para_docx(conteudo, docx_path)
            word = _caminho_word()
            if word:
                subprocess.Popen([word, str(docx_path)])  # nosec B603
            else:
                os.startfile(str(docx_path))  # nosec B606
        elif md_path.exists():
            os.startfile(str(md_path))  # nosec B606
    if ajax:
        return {"ok": True}
    return RedirectResponse("/", status_code=303)


# ── editar e registrar revisão ───────────────────────────────────────────────

@app.get("/editar/{doc_id}", response_class=HTMLResponse)
def editar_documento(doc_id: int, request: Request, aviso: str = "", db: Session = Depends(get_db)):
    doc = db.query(Documento).filter(Documento.id == doc_id).first()
    if not doc:
        return HTMLResponse("Documento não encontrado.", status_code=404)
    tem_arquivo = bool(doc.arquivo and doc.categoria != "Certificado")
    conteudo_html = _renderizar(doc.arquivo) if tem_arquivo else ""
    tem_word_pendente = (WORD_TEMP_DIR / f"{doc_id}.docx").exists()
    return templates.TemplateResponse(request, "editar.html", {
        "request":           request,
        "doc":               doc,
        "tem_arquivo":       tem_arquivo,
        "conteudo_html":     conteudo_html,
        "tem_word_pendente": tem_word_pendente,
        "aviso":             aviso,
    })


@app.get("/descartar-word/{doc_id}")
def descartar_word(doc_id: int):
    docx_temp = WORD_TEMP_DIR / f"{doc_id}.docx"
    try:
        docx_temp.unlink(missing_ok=True)
    except PermissionError:
        return RedirectResponse(f"/editar/{doc_id}?aviso=feche-o-word", status_code=303)
    return RedirectResponse(f"/editar/{doc_id}", status_code=303)


@app.post("/editar/{doc_id}")
def salvar_revisao(
    doc_id:  int,
    notas:   str = Form(""),
    revisor: str = Form("Michel Rui Costa"),
    db: Session = Depends(get_db),
):
    doc = db.query(Documento).filter(Documento.id == doc_id).first()
    if doc:
        # Se o usuário editou no Word, converte o .docx de volta para .md
        docx_temp = WORD_TEMP_DIR / f"{doc_id}.docx"
        if docx_temp.exists() and doc.arquivo and doc.arquivo.endswith(".md"):
            try:
                md_novo = _docx_para_md(docx_temp)
                (FINDABC_DIR / doc.arquivo).write_text(md_novo, encoding="utf-8")
            except Exception:
                pass  # nosec B110 — não bloqueia o registro da revisão
            finally:
                docx_temp.unlink(missing_ok=True)

        agora = datetime.now()
        db.add(Revisao(
            doc_id=doc_id,
            versao=doc.proxima_versao_sugerida,
            data=agora.date(),
            hora=agora.strftime("%H:%M"),
            responsavel=revisor,
            notas=notas or "—",
        ))
        db.commit()
    return RedirectResponse("/", status_code=303)


# ── pacote de due diligence ──────────────────────────────────────────────────

@app.get("/pacote", response_class=HTMLResponse)
def pacote_due_diligence(request: Request, db: Session = Depends(get_db)):
    """Opção 1: página imprimível com índice de todos os documentos próprios."""
    docs = (
        db.query(Documento)
        .filter(Documento.categoria != "Certificado")
        .order_by(Documento.categoria, Documento.nome)
        .all()
    )
    return templates.TemplateResponse(request, "pacote.html", {
        "request": request,
        "docs":    docs,
        "hoje":    date.today(),
    })

# ── segurança — painel de auditoria de segurança ─────────────────────────────

@app.get("/seguranca", response_class=HTMLResponse)
def seguranca(request: Request):
    summary_path = AUDIT_REPORT_DIR / "latest_summary.json"
    history_path = AUDIT_REPORT_DIR / "history.json"
    summary = None
    history = []
    if summary_path.exists():
        try:
            summary = json.loads(summary_path.read_text(encoding="utf-8"))
        except Exception:
            pass
    if history_path.exists():
        try:
            history = json.loads(history_path.read_text(encoding="utf-8"))
        except Exception:
            pass
    with _audit_lock:
        audit_st = _audit_status.copy()
    return templates.TemplateResponse(request, "seguranca.html", {
        "request":    request,
        "summary":    summary,
        "history":    history,
        "audit_st":   audit_st,
    })


@app.post("/seguranca/rodar")
def seguranca_rodar():
    with _audit_lock:
        if _audit_status["running"]:
            return {"ok": False, "message": "Auditoria já em andamento"}
        _audit_status["running"]     = True
        _audit_status["started_at"]  = datetime.now().strftime("%d/%m/%Y %H:%M")
        _audit_status["finished_at"] = None
        _audit_status["error"]       = None
    t = threading.Thread(target=_run_audit_bg, daemon=True)
    t.start()
    return {"ok": True, "message": "Auditoria iniciada"}


@app.get("/seguranca/status")
def seguranca_status():
    with _audit_lock:
        return _audit_status.copy()


@app.get("/seguranca/exportar")
def seguranca_exportar():
    summary_path = AUDIT_REPORT_DIR / "latest_summary.json"
    history_path = AUDIT_REPORT_DIR / "history.json"
    if not summary_path.exists():
        return Response("Nenhuma auditoria disponível.", status_code=404)
    try:
        xlsx = gerar_excel(summary_path, history_path, LOGO_PATH if LOGO_PATH.exists() else None)
    except Exception as e:
        return Response(f"Erro ao gerar Excel: {e}", status_code=500)
    ts = datetime.now().strftime("%Y-%m-%d")
    filename = f"auditoria_seguranca_{ts}.xlsx"
    return Response(
        content=xlsx,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── /p/{token} — reservado para Opção 2 (portal do cliente, pós-deploy) ─────
# @app.get("/p/{token}", response_class=HTMLResponse)
# def portal_cliente(token: str, ...): ...
